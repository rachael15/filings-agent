"""
Filings -> summary + small working model.

    python filings_agent.py AAPL
    python filings_agent.py AAPL --transcript call.txt --out out

Pipeline
  1. SEC EDGAR: ticker -> CIK, XBRL "company facts" (structured financials),
     latest 10-K/10-Q text, latest earnings-release 8-K (Item 2.02, Ex. 99.1).
  2. Claude: written summary + proposed bear/base/bull assumptions (JSON).
  3. openpyxl: Excel model with live formulas and a scenario switch.

Env: ANTHROPIC_API_KEY, SEC_USER_AGENT ("Your Name you@example.com" - SEC requires it).
"""
import argparse
import json
import os
import re
import sys
from datetime import date
from pathlib import Path

import requests
from anthropic import Anthropic
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

MODEL = os.environ.get("CLAUDE_MODEL", "claude-sonnet-5-5")
UA = os.environ.get("SEC_USER_AGENT", "")
N_HIST = 3

# XBRL concepts, in priority order; later ones only fill years the earlier ones lack.
CONCEPTS = {
    "revenue": (["Revenues", "RevenueFromContractWithCustomerExcludingAssessedTax", "SalesRevenueNet"], "USD"),
    "op_income": (["OperatingIncomeLoss"], "USD"),
    "net_income": (["NetIncomeLoss", "ProfitLoss"], "USD"),
    "shares": (["WeightedAverageNumberOfDilutedSharesOutstanding"], "shares"),
    "da": (["DepreciationDepletionAndAmortization", "DepreciationAndAmortization",
            "DepreciationAmortizationAndAccretionNet"], "USD"),
    "capex": (["PaymentsToAcquirePropertyPlantAndEquipment", "PaymentsToAcquireProductiveAssets"], "USD"),
}


# ---------------------------------------------------------------- EDGAR ----
def sec_get(url, **kw):
    if not UA:
        sys.exit('Set SEC_USER_AGENT, e.g. "Jane Doe jane@example.com" (SEC blocks anonymous clients).')
    r = requests.get(url, headers={"User-Agent": UA}, timeout=60, **kw)
    r.raise_for_status()
    return r


def ticker_to_cik(ticker):
    data = sec_get("https://www.sec.gov/files/company_tickers.json").json()
    for row in data.values():
        if row["ticker"].upper() == ticker.upper():
            return int(row["cik_str"]), row["title"]
    sys.exit(f"Ticker {ticker} not found on EDGAR")


def annual_series(facts, names, unit):
    """{period_end_iso: value} for full-year figures from 10-Ks, merged across concepts."""
    out = {}
    gaap = facts.get("facts", {}).get("us-gaap", {})
    for name in names:
        best = {}
        for it in gaap.get(name, {}).get("units", {}).get(unit, []):
            if it.get("form") != "10-K" or it.get("fp") != "FY":
                continue
            if "start" in it:
                days = (date.fromisoformat(it["end"]) - date.fromisoformat(it["start"])).days
                if not 340 <= days <= 380:
                    continue
            cur = best.get(it["end"])
            if cur is None or it["filed"] > cur["filed"]:
                best[it["end"]] = it
        for end, it in best.items():
            if not any(abs((date.fromisoformat(end) - date.fromisoformat(e)).days) <= 10 for e in out):
                out[end] = it["val"]
    return out


def lookup(series, end):
    for e, v in series.items():
        if abs((date.fromisoformat(e) - date.fromisoformat(end)).days) <= 10:
            return v
    return None


def build_financials(facts):
    series = {k: annual_series(facts, n, u) for k, (n, u) in CONCEPTS.items()}
    ends = sorted(series["revenue"])[-N_HIST:]
    if not ends:
        sys.exit("No annual revenue found in XBRL data (non-US-GAAP filer?).")
    scale = {"shares": 1e6}
    rows = {}
    for k, s in series.items():
        rows[k] = [None if (v := lookup(s, e)) is None else v / scale.get(k, 1e6) for e in ends]
    return {"years": [e[:4] for e in ends], "ends": ends, "rows": rows}  # $ and shares in millions


def filing_url(cik, acc, doc):
    return f"https://www.sec.gov/Archives/edgar/data/{cik}/{acc.replace('-', '')}/{doc}"


def html_to_text(html):
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style"]):
        t.decompose()
    return re.sub(r"\n\s*\n+", "\n\n", soup.get_text("\n"))


def pick_sections(text, mda_chars=70_000, risk_chars=20_000):
    """Skip the table of contents by taking the later 'MD&A' / 'Risk Factors' hit."""
    def grab(pattern, n):
        hits = [m.start() for m in re.finditer(pattern, text, re.I)]
        if not hits:
            return ""
        return text[hits[1] if len(hits) > 1 else hits[0]:][:n]
    return (grab(r"management.s discussion and analysis", mda_chars)
            + "\n\n--- RISK FACTORS ---\n\n" + grab(r"item\s+1a\.?\s+risk factors", risk_chars))


def latest_documents(cik, subs):
    rec = subs["filings"]["recent"]
    n = len(rec["form"])
    periodic = earnings = None
    for i in range(n):
        form = rec["form"][i]
        if periodic is None and form in ("10-K", "10-Q"):
            periodic = (form, rec["filingDate"][i], filing_url(cik, rec["accessionNumber"][i], rec["primaryDocument"][i]))
        if earnings is None and form == "8-K" and "2.02" in rec["items"][i]:
            earnings = (rec["filingDate"][i], rec["accessionNumber"][i])
        if periodic and earnings:
            break
    out = {}
    if periodic:
        form, filed, url = periodic
        out["periodic"] = {"form": form, "filed": filed, "url": url,
                           "text": pick_sections(html_to_text(sec_get(url).text))}
    if earnings:
        filed, acc = earnings
        idx = sec_get(f"https://www.sec.gov/Archives/edgar/data/{cik}/{acc.replace('-', '')}/index.json").json()
        names = [f["name"] for f in idx["directory"]["item"]]
        pr = next((n for n in names if re.search(r"ex-?99[-_.]?1", n, re.I) and n.endswith((".htm", ".html"))), None)
        if pr:
            url = filing_url(cik, acc, pr)
            out["release"] = {"filed": filed, "url": url, "text": html_to_text(sec_get(url).text)[:40_000]}
    return out


# --------------------------------------------------------------- Claude ----
def ask(client, system, user, max_tokens):
    msg = client.messages.create(model=MODEL, max_tokens=max_tokens, system=system,
                                 messages=[{"role": "user", "content": user}])
    return "".join(b.text for b in msg.content if b.type == "text")


def summarize(client, name, fin, docs, transcript):
    parts = [f"Company: {name}", "Annual financials ($M, shares M):\n" + json.dumps(fin, indent=1)]
    if "release" in docs:
        parts.append(f"EARNINGS RELEASE ({docs['release']['filed']}):\n{docs['release']['text']}")
    if "periodic" in docs:
        p = docs["periodic"]
        parts.append(f"{p['form']} ({p['filed']}) excerpts:\n{p['text']}")
    if transcript:
        parts.append(f"EARNINGS CALL TRANSCRIPT:\n{transcript[:60_000]}")
    system = ("You are an equity research analyst. Use only the provided material; cite the source "
              "(10-K/10-Q, release, transcript) for each figure. If something isn't in the material, say so.")
    task = ("Write a Markdown note with: 1) Headline takeaways (5 bullets) 2) Results vs. prior period "
            "3) Segment / driver commentary 4) Guidance and management tone (quote sparingly) "
            "5) Key risks and what changed 6) Questions to dig into. Under 700 words.")
    return ask(client, system, "\n\n".join(parts) + "\n\n" + task, 3000)


def propose_assumptions(client, summary, fin):
    prompt = f"""Based on this note and the history, propose 3-year projection assumptions.
Return ONLY JSON, no prose, exactly this shape (decimals, e.g. 0.05 = 5%):
{{"revenue_growth": {{"bear": [y1,y2,y3], "base": [y1,y2,y3], "bull": [y1,y2,y3]}},
 "operating_margin": {{"bear": [y1,y2,y3], "base": [y1,y2,y3], "bull": [y1,y2,y3]}},
 "tax_rate": 0.21, "da_pct_revenue": 0.0, "capex_pct_revenue": 0.0, "share_change": 0.0,
 "rationale": "2-3 sentences"}}

NOTE:
{summary}

HISTORY:
{json.dumps(fin)}"""
    raw = ask(client, "You output strict JSON only.", prompt, 1200)
    try:
        return json.loads(re.search(r"\{.*\}", raw, re.S).group(0))
    except Exception:
        return None


def fallback_assumptions(fin):
    rev = [v for v in fin["rows"]["revenue"] if v]
    g = (rev[-1] / rev[0]) ** (1 / (len(rev) - 1)) - 1 if len(rev) > 1 else 0.05
    oi, r = fin["rows"]["op_income"][-1], fin["rows"]["revenue"][-1]
    m = oi / r if oi and r else 0.15
    return {"revenue_growth": {"bear": [g - .04] * 3, "base": [g] * 3, "bull": [g + .04] * 3},
            "operating_margin": {"bear": [m - .02] * 3, "base": [m] * 3, "bull": [m + .02] * 3},
            "tax_rate": 0.21, "da_pct_revenue": 0.04, "capex_pct_revenue": 0.04, "share_change": 0.0,
            "rationale": "Fallback: history-based CAGR/margin +/- spread (LLM assumptions unavailable)."}


# ---------------------------------------------------------------- Excel ----
BLUE = Font(color="0000FF")        # inputs
BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="DDE6F0")


def build_workbook(path, name, fin, a):
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"], ws["A1"].font = f"{name} - simplified operating model ($M except per share)", Font(bold=True, size=13)
    ws["A3"], ws["B3"] = "Scenario (1=Bear, 2=Base, 3=Bull)", 2
    ws["B3"].font = BLUE
    ws["A4"] = "Blue = input. Net income = EBIT x (1 - tax); ignores interest/other. Not investment advice."

    # Assumptions: columns E:G = projection years 1-3
    yr = int(fin["years"][-1])
    for c, y in zip("EFG", range(yr + 1, yr + 4)):
        ws[f"{c}5"] = f"{y}E"
        ws[f"{c}5"].font = BOLD
    ws["A5"], ws["A5"].font = "Assumptions", BOLD

    def block(row, label, key):
        for i, sc in enumerate(("bear", "base", "bull")):
            ws[f"A{row + i}"] = f"{label} - {sc.title()}"
            for c, v in zip("EFG", a[key][sc]):
                ws[f"{c}{row + i}"], ws[f"{c}{row + i}"].font = v, BLUE
                ws[f"{c}{row + i}"].number_format = "0.0%"
        ws[f"A{row + 3}"] = f"{label} - Selected"
        for c in "EFG":
            ws[f"{c}{row + 3}"] = f"=CHOOSE($B$3,{c}{row},{c}{row + 1},{c}{row + 2})"
            ws[f"{c}{row + 3}"].number_format = "0.0%"
    block(6, "Revenue growth", "revenue_growth")      # selected row 9
    block(11, "Operating margin", "operating_margin")  # selected row 14
    for row, label, v in ((16, "Tax rate", a["tax_rate"]), (17, "D&A % revenue", a["da_pct_revenue"]),
                          (18, "Capex % revenue", a["capex_pct_revenue"]), (19, "Diluted share change", a["share_change"])):
        ws[f"A{row}"] = label
        for c in "EFG":
            ws[f"{c}{row}"], ws[f"{c}{row}"].font = v, BLUE
            ws[f"{c}{row}"].number_format = "0.0%"

    # Projection
    ws["A22"] = "Line item"
    for c, y in zip("BCD", fin["years"]):
        ws[f"{c}22"] = f"{y}A"
    for c in "EFG":
        ws[f"{c}22"] = f"={c}5"
    for c in "ABCDEFG":
        ws[f"{c}22"].font, ws[f"{c}22"].fill = BOLD, HDR

    labels = {23: "Revenue", 24: "  growth %", 25: "Operating income", 26: "  margin %", 27: "Taxes",
              28: "Net income", 29: "Diluted shares (M)", 30: "EPS (diluted)", 31: "D&A", 32: "Capex",
              33: "Free cash flow (NI + D&A - capex)"}
    for r, t in labels.items():
        ws[f"A{r}"] = t
    hist = {23: "revenue", 25: "op_income", 28: "net_income", 29: "shares", 31: "da", 32: "capex"}
    for c, i in zip("BCD", range(N_HIST)):
        for r, k in hist.items():
            v = fin["rows"][k][i] if i < len(fin["rows"][k]) else None
            if v is not None:
                ws[f"{c}{r}"] = abs(v) if k == "capex" else v
    cols = "BCDEFG"
    for j, c in enumerate(cols):
        p = cols[j - 1] if j else None
        proj = c in "EFG"
        if proj:
            ws[f"{c}23"] = f"={p}23*(1+{c}9)"
            ws[f"{c}25"] = f"={c}23*{c}14"
            ws[f"{c}27"] = f"={c}25*{c}16"
            ws[f"{c}28"] = f"={c}25-{c}27"
            ws[f"{c}29"] = f"={p}29*(1+{c}19)"
            ws[f"{c}31"] = f"={c}23*{c}17"
            ws[f"{c}32"] = f"={c}23*{c}18"
        if p:
            ws[f"{c}24"] = f"=IFERROR({c}23/{p}23-1,\"\")"
        ws[f"{c}26"] = f"=IFERROR({c}25/{c}23,\"\")"
        ws[f"{c}30"] = f"=IFERROR({c}28/{c}29,\"\")"
        ws[f"{c}33"] = f"=IFERROR({c}28+{c}31-{c}32,\"\")"
        for r in (23, 25, 27, 28, 29, 31, 32, 33):
            ws[f"{c}{r}"].number_format = "#,##0;(#,##0)"
        for r in (24, 26):
            ws[f"{c}{r}"].number_format = "0.0%"
        ws[f"{c}30"].number_format = "0.00"
    ws["A35"] = "Rationale"
    ws["A35"].font = BOLD
    ws["A36"] = a.get("rationale", "")
    ws.column_dimensions["A"].width = 38
    for c in "BCDEFG":
        ws.column_dimensions[c].width = 12
    wb.save(path)


# ----------------------------------------------------------------- main ----
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("ticker")
    ap.add_argument("--transcript", help="path to a transcript .txt (not available on EDGAR)")
    ap.add_argument("--out", default="out")
    ap.add_argument("--mode", choices=["api", "dump", "build"], default="api",
                    help="api: call Claude API; dump: save fetched data only; build: use saved summary/assumptions")
    args = ap.parse_args()

    out = Path(args.out)
    out.mkdir(exist_ok=True)
    t = args.ticker.upper()

    cik, name = ticker_to_cik(t)
    print(f"{name} (CIK {cik})")
    facts = sec_get(f"https://data.sec.gov/api/xbrl/companyfacts/CIK{cik:010d}.json").json()
    subs = sec_get(f"https://data.sec.gov/submissions/CIK{cik:010d}.json").json()
    fin = build_financials(facts)
    docs = latest_documents(cik, subs)
    print("Fetched:", ", ".join(f"{k} ({v['filed']})" for k, v in docs.items()) or "no filings text")
    transcript = Path(args.transcript).read_text(encoding="utf-8") if args.transcript else None

    if args.mode == "dump":  # no API key: save context for Claude Code to read and summarize
        (out / f"{t}_context.json").write_text(json.dumps(
            {"name": name, "financials": fin, "docs": docs, "transcript": transcript}, indent=1), encoding="utf-8")
        print(f"Wrote {out / (t + '_context.json')}")
        return
    if args.mode == "build":  # summary + assumptions were written by hand / by Claude Code
        summary = (out / f"{t}_summary_body.md").read_text(encoding="utf-8")
        assumptions = json.loads((out / f"{t}_assumptions.json").read_text(encoding="utf-8"))
    else:
        client = Anthropic()
        summary = summarize(client, name, fin, docs, transcript)
        assumptions = propose_assumptions(client, summary, fin) or fallback_assumptions(fin)

    sources = [f"- {v.get('form', 'Earnings release 8-K')}: {v['url']}" for v in docs.values()]
    (out / f"{t}_summary.md").write_text(f"# {name} ({t})\n\n{summary}\n\n## Sources\n" + "\n".join(sources) + "\n",
                                         encoding="utf-8")
    build_workbook(out / f"{t}_model.xlsx", name, fin, assumptions)
    print(f"Wrote {out / (t + '_summary.md')} and {out / (t + '_model.xlsx')}")


if __name__ == "__main__":
    main()
